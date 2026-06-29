"""
Generates cogmodel_presentation.pptx — self-presentation for the MIPT
Center for Cognitive Modeling master's program interview.

Style: minimalist, light background, no bright colors. One muted slate accent.
~5 minute talk, 5 slides.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUTPUT = "/home/chikibriki/LoqicQA/cogmodel_presentation.pptx"

# ---------------------------------------------------------------------------
# Palette — light, muted, no bright colors
# ---------------------------------------------------------------------------
C_BG       = RGBColor(0xFB, 0xFB, 0xF9)   # warm off-white background
C_INK      = RGBColor(0x2B, 0x2B, 0x2B)   # main near-black text
C_GRAY     = RGBColor(0x6E, 0x6E, 0x6E)   # secondary gray text
C_FAINT    = RGBColor(0x9A, 0x9A, 0x9A)   # faint gray (slide no., captions)
C_ACCENT   = RGBColor(0x55, 0x68, 0x7A)   # muted slate-blue accent
C_RULE     = RGBColor(0xD8, 0xD8, 0xD2)   # hairline rule
C_CHIP     = RGBColor(0xEF, 0xEF, 0xEA)   # subtle chip / panel fill

FONT = "Calibri"

W = Inches(13.333)
H = Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_masters[0].slide_layouts[6]

MARGIN = Inches(0.9)
CONTENT_W = W - 2 * MARGIN

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def set_bg(slide, color=C_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def style_run(run, size, color=C_INK, bold=False, italic=False, font=FONT, spacing=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def set_para(p, align=PP_ALIGN.LEFT, space_after=6, space_before=0, line=None):
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line


def add_line(slide, left, top, width, color=C_RULE, weight=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(weight))
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln


def add_panel(slide, left, top, width, height, color=C_CHIP):
    from pptx.enum.shapes import MSO_SHAPE
    p = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    p.fill.solid()
    p.fill.fore_color.rgb = color
    p.line.fill.background()
    p.shadow.inherit = False
    try:
        p.adjustments[0] = 0.04
    except Exception:
        pass
    return p


def slide_header(slide, number, kicker, title):
    """Slide number (top-right), small kicker, title, accent rule."""
    set_bg(slide)
    # slide number
    _, tf = textbox(slide, W - Inches(1.4), Inches(0.45), Inches(0.8), Inches(0.4))
    p = tf.paragraphs[0]
    set_para(p, align=PP_ALIGN.RIGHT, space_after=0)
    r = p.add_run(); r.text = f"{number:02d}"
    style_run(r, 12, color=C_FAINT)

    # kicker
    _, tf = textbox(slide, MARGIN, Inches(0.6), CONTENT_W, Inches(0.35))
    p = tf.paragraphs[0]
    set_para(p, space_after=0)
    r = p.add_run(); r.text = kicker.upper()
    style_run(r, 11.5, color=C_ACCENT, bold=True)
    r.font._rPr.set("spc", "180")  # letter spacing

    # title
    _, tf = textbox(slide, MARGIN, Inches(0.92), CONTENT_W, Inches(0.7))
    p = tf.paragraphs[0]
    set_para(p, space_after=0)
    r = p.add_run(); r.text = title
    style_run(r, 30, color=C_INK, bold=True)

    # accent rule (short)
    add_line(slide, MARGIN, Inches(1.62), Inches(1.1), color=C_ACCENT, weight=2.2)


def bullet(tf, text, size=16, color=C_INK, bold=False, level=0,
           space_after=8, first=False, marker="—", italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    set_para(p, space_after=space_after, line=1.08)
    p.level = level
    indent = "      " * level
    if marker:
        rm = p.add_run(); rm.text = f"{indent}{marker}  "
        style_run(rm, size, color=C_ACCENT, bold=False)
    elif indent:
        rm = p.add_run(); rm.text = indent
        style_run(rm, size, color=color)
    r = p.add_run(); r.text = text
    style_run(r, size, color=color, bold=bold, italic=italic)
    return p


def add_chips(slide, left, top, items, gap=Inches(0.18), size=12):
    """Render a row of subtle tag chips."""
    x = left
    for it in items:
        w = Inches(0.22 + 0.105 * len(it))
        panel = add_panel(slide, x, top, w, Inches(0.42), color=C_CHIP)
        tf = panel.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
        tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; set_para(p, align=PP_ALIGN.CENTER, space_after=0)
        r = p.add_run(); r.text = it
        style_run(r, size, color=C_GRAY)
        x = x + w + gap

# ===========================================================================
# SLIDE 1 — Title / About
# ===========================================================================
s = prs.slides.add_slide(BLANK)
set_bg(s)

# left text block
_, tf = textbox(s, MARGIN, Inches(2.35), Inches(8.2), Inches(3.0))
p = tf.paragraphs[0]; set_para(p, space_after=2)
r = p.add_run(); r.text = "Сахаров Даниэль Александрович"
style_run(r, 38, color=C_INK, bold=True)

p = tf.add_paragraph(); set_para(p, space_after=14, space_before=2)
r = p.add_run(); r.text = "Бакалавр ВШПИ МФТИ, 2026"
style_run(r, 18, color=C_GRAY)

p = tf.add_paragraph(); set_para(p, space_after=0, line=1.15)
r = p.add_run(); r.text = "Собеседование в магистратуру"
style_run(r, 16, color=C_ACCENT, bold=True)
p = tf.add_paragraph(); set_para(p, space_after=0, line=1.15)
r = p.add_run(); r.text = "Центр когнитивного моделирования МФТИ"
style_run(r, 16, color=C_ACCENT, bold=True)

# accent rule
add_line(s, MARGIN, Inches(2.2), Inches(1.1), color=C_ACCENT, weight=2.4)

# github line
_, tf = textbox(s, MARGIN, Inches(5.35), Inches(6.0), Inches(0.4))
p = tf.paragraphs[0]; set_para(p, space_after=0)
r = p.add_run(); r.text = "github.com/grgtr"
style_run(r, 15, color=C_FAINT)

# photo placeholder panel (right)
ph = add_panel(s, W - Inches(4.3), Inches(2.1), Inches(3.0), Inches(3.4), color=C_CHIP)
tf = ph.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; set_para(p, align=PP_ALIGN.CENTER, space_after=0)
r = p.add_run(); r.text = "фото"
style_run(r, 14, color=C_FAINT)

# ===========================================================================
# SLIDE 2 — Achievements
# ===========================================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, 2, "Достижения", "Чем я горжусь")

col_w = Inches(5.55)
top = Inches(2.05)

# Left column — competitions
_, tf = textbox(s, MARGIN, top, col_w, Inches(4.6))
p = tf.paragraphs[0]; set_para(p, space_after=10)
r = p.add_run(); r.text = "Соревнования и олимпиады"
style_run(r, 15, color=C_ACCENT, bold=True)
bullet(tf, "Призёр Deep Learning School, 2026", size=15)
bullet(tf, "1 место — командный хакатон VK, улучшение ответов Маруси, 2025", size=15)
bullet(tf, "Участник AIDAO", size=15)
bullet(tf, "Призёр олимпиад Физтеха по математике и физике", size=15)

# Right column — science
rx = MARGIN + col_w + Inches(0.7)
_, tf = textbox(s, rx, top, col_w, Inches(2.7))
p = tf.paragraphs[0]; set_para(p, space_after=10)
r = p.add_run(); r.text = "Научные работы · 68-я конференция МФТИ"
style_run(r, 15, color=C_ACCENT, bold=True)
bullet(tf, "Применение сигнатурных представлений многомерных путей в анализе временных рядов", size=14)
bullet(tf, "Система оценки качества моделей детекции логических аномалий", size=14)

# highlight panel — diploma result
hp = add_panel(s, rx, Inches(4.55), col_w, Inches(1.55), color=C_CHIP)
tf = hp.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
tf.margin_top = Inches(0.18); tf.margin_bottom = Inches(0.18)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; set_para(p, space_after=4)
r = p.add_run(); r.text = "Результат диплома"
style_run(r, 13, color=C_ACCENT, bold=True)
p = tf.add_paragraph(); set_para(p, space_after=0, line=1.1)
r = p.add_run(); r.text = "Детекция логических аномалий: уровень GPT-4o по AUROC на более слабой открытой VLM — InternVL-8B"
style_run(r, 14.5, color=C_INK, bold=True)

# ===========================================================================
# SLIDE 3 — Why the program + test task
# ===========================================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, 3, "Мотивация", "Почему ЦКМ и направление VLA")

# path chips
_, tf = textbox(s, MARGIN, Inches(1.95), CONTENT_W, Inches(0.4))
p = tf.paragraphs[0]; set_para(p, space_after=0)
r = p.add_run(); r.text = "Путь:  "
style_run(r, 14, color=C_GRAY, bold=True)
r = p.add_run(); r.text = "Computer Vision (ViT · MAE · DINO · DETR · CLIP)  →  диплом по CV  →  World Models + RL  →  VLA"
style_run(r, 14, color=C_INK)

# panel — test task
panel = add_panel(s, MARGIN, Inches(2.7), CONTENT_W, Inches(2.55), color=C_CHIP)
tf = panel.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3)
tf.margin_top = Inches(0.22); tf.margin_bottom = Inches(0.2)
p = tf.paragraphs[0]; set_para(p, space_after=10)
r = p.add_run(); r.text = "Выполнил тестовое задание на стажёрскую позицию по направлению"
style_run(r, 16, color=C_ACCENT, bold=True)
bullet(tf, "Сравнил explicit (DreamerV3 / RSSM) и implicit (TD-MPC2) модели мира на ManiSkill2 LiftCube и DMControl", size=14.5)
bullet(tf, "Объяснил коллапс KL-утилизации у explicit-агента через linear probing R² и CCA латентов", size=14.5)
bullet(tf, "Дообучил VLA-модель SmolVLA (~450M, SigLIP) через LoRA; показал, что её представления комплементарны латентам моделей мира", size=14.5)

# closing line
_, tf = textbox(s, MARGIN, Inches(5.55), CONTENT_W, Inches(1.4))
p = tf.paragraphs[0]; set_para(p, space_after=8, line=1.12)
r = p.add_run(); r.text = "Поэтому мне особенно близка Лаборатория воплощённого ИИ и треки по VLA."
style_run(r, 15.5, color=C_INK, bold=True)
p = tf.add_paragraph(); set_para(p, space_after=0, line=1.1)
r = p.add_run(); r.text = "Интересные курсы программы:  "
style_run(r, 14, color=C_GRAY, bold=True)
r = p.add_run(); r.text = "[вписать 2–3 названия с сайта]"
style_run(r, 14, color=C_FAINT, italic=True)

# ===========================================================================
# SLIDE 4 — Research interests
# ===========================================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, 4, "Научные интересы", "Тема исследований")

# the topic statement
_, tf = textbox(s, MARGIN, Inches(1.95), CONTENT_W, Inches(0.9))
p = tf.paragraphs[0]; set_para(p, space_after=0, line=1.1)
r = p.add_run(); r.text = "Vision-language-action модели для робототехники: "
style_run(r, 21, color=C_INK, bold=True)
r = p.add_run(); r.text = "память и планирование действий"
style_run(r, 21, color=C_ACCENT, bold=True)

# direction chips
add_chips(s, MARGIN, Inches(2.95),
          ["VLA-модели", "модели мира", "память в RL", "обобщаемость в RL", "мобильная манипуляция"],
          size=12.5)

# reference works
_, tf = textbox(s, MARGIN, Inches(3.75), CONTENT_W, Inches(2.4))
p = tf.paragraphs[0]; set_para(p, space_after=10)
r = p.add_run(); r.text = "Опорные работы и связь с лабораторией"
style_run(r, 15, color=C_ACCENT, bold=True)
bullet(tf, "LERa (Look, Explain, Replan) — реплэннинг по визуальной обратной связи · группа А. Ковалёва, AIRI", size=15)
bullet(tf, "ELMUR, MIKASA — внешняя память агентов и бенчмарк памяти роботов · ЦКМ / AIRI", size=15)
bullet(tf, "VLA с визуальным промптингом — проекция трасс ключевых точек на карты глубины", size=15)

# bridge line
_, tf = textbox(s, MARGIN, Inches(6.25), CONTENT_W, Inches(0.9))
p = tf.paragraphs[0]; set_para(p, space_after=0, line=1.12)
r = p.add_run(); r.text = "Мостик: в дипломе я строил систему метрик для оценки рассуждений VLM — умею не только улучшать модель, но и измерять, где она ошибается."
style_run(r, 14, color=C_GRAY, italic=True)

# ===========================================================================
# SLIDE 5 — Future image
# ===========================================================================
s = prs.slides.add_slide(BLANK)
slide_header(s, 5, "Образ будущего", "Через 3–4 года")

# timeline
ty = Inches(2.7)
steps = ["Магистратура", "Первые публикации\nи конференции", "Аспирантура /\nR&D-исследователь"]
seg_w = Inches(3.6)
gap = Inches(0.5)
x = MARGIN
add_line(s, MARGIN + Inches(0.3), ty + Inches(0.55), CONTENT_W - Inches(0.6),
         color=C_RULE, weight=1.4)
for i, st in enumerate(steps):
    panel = add_panel(s, x, ty, seg_w, Inches(1.15), color=C_CHIP)
    tf = panel.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    for j, linetext in enumerate(st.split("\n")):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        set_para(p, align=PP_ALIGN.CENTER, space_after=0, line=1.05)
        r = p.add_run(); r.text = linetext
        style_run(r, 15, color=C_INK, bold=True)
    x = x + seg_w + gap

# anchors
_, tf = textbox(s, MARGIN, Inches(4.55), CONTENT_W, Inches(0.5))
p = tf.paragraphs[0]; set_para(p, align=PP_ALIGN.CENTER, space_after=0)
r = p.add_run(); r.text = "изучать новое   ·   создавать новое   ·   защищать идеи"
style_run(r, 17, color=C_ACCENT, bold=True)

# closing sentence
_, tf = textbox(s, MARGIN + Inches(1.0), Inches(5.5), CONTENT_W - Inches(2.0), Inches(1.2))
p = tf.paragraphs[0]; set_para(p, align=PP_ALIGN.CENTER, space_after=0, line=1.18)
r = p.add_run(); r.text = "Исследователь в области обучения роботов и VLA-моделей — формулирую задачи, строю решения и выношу их на конференции."
style_run(r, 15, color=C_GRAY)

# ---------------------------------------------------------------------------
prs.save(OUTPUT)
print("Saved:", OUTPUT, "| slides:", len(prs.slides._sldIdLst))
